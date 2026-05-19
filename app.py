import os
import json
from datetime import datetime
from dotenv import load_dotenv

# Load environment variables FIRST (before any module that uses them)
load_dotenv()

# Then import Flask and other dependencies
from flask import (
    Flask, render_template, request, redirect,
    url_for, flash, session, get_flashed_messages
)
from functools import wraps
from google.cloud import secretmanager
from werkzeug.security import generate_password_hash

# Now import modules that may rely on environment variables
from email_notify import (
    notify_supervisor_topic_submitted,
    notify_supervisor_topic_dropped,
    notify_student_cleared,
    notify_coordinator_cleared,
    notify_coordinator_proposal_submitted  
)

from sheet import get_supervisor_email, get_student_email, get_student_name

app = Flask(__name__)

# ── SECRET_KEY ────────────────────────────────────────────────────────────────
def get_secret_key():
    if os.getenv('GAE_ENV', '').startswith('standard'):
        project_id = os.getenv('GOOGLE_CLOUD_PROJECT')
        client = secretmanager.SecretManagerServiceClient()
        name = f"projects/{project_id}/secrets/SECRET_KEY/versions/latest"
        resp = client.access_secret_version(request={'name': name})
        return resp.payload.data.decode('UTF-8')
    key = os.getenv('SECRET_KEY')
    if not key:
        raise RuntimeError('SECRET_KEY not set')
    return key

app.secret_key = get_secret_key()

# ── Google Sheets credentials ─────────────────────────────────────────────────
def get_google_credentials():
    if os.getenv('GAE_ENV', '').startswith('standard'):
        project_id = os.getenv('GOOGLE_CLOUD_PROJECT')
        client = secretmanager.SecretManagerServiceClient()
        name = f"projects/{project_id}/secrets/google-credentials-json/versions/latest"
        resp = client.access_secret_version(request={'name': name})
        return resp.payload.data.decode('UTF-8')
    return os.getenv('GOOGLE_CREDENTIALS_JSON')

# ── Passphrases ───────────────────────────────────────────────────────────────
def _get_env(var):
    return (os.getenv(var) or '').strip()

def _get_gcp_secret(secret_name):
    try:
        project = os.getenv('GOOGLE_CLOUD_PROJECT')
        client = secretmanager.SecretManagerServiceClient()
        name = f"projects/{project}/secrets/{secret_name}/versions/latest"
        resp = client.access_secret_version(request={"name": name})
        return resp.payload.data.decode("UTF-8").strip()
    except Exception:
        return ''

def get_coordinator_passphrase():
    if os.getenv('GAE_ENV', '').startswith('standard'):
        return _get_gcp_secret('coordinator-passphrase') or _get_gcp_secret('staff-passphrase')
    return _get_env('COORDINATOR_PASSPHRASE') or _get_env('STAFF_PASSPHRASE')

def get_supervisor_passphrase():
    if os.getenv('GAE_ENV', '').startswith('standard'):
        return _get_gcp_secret('staff-passphrase')
    return _get_env('STAFF_PASSPHRASE')

# ── Business logic imports ────────────────────────────────────────────────────
import sheet
from sheet import (
    CURRENT_SESSION,
    PREVIOUS_SESSION,
    backfill_sessions,
    update_student_session,
    edit_student_record,
    backfill_submission_dates,
    submit_topic_proposal,
    get_presentation_records,
    supervisor_clear_student,
    coordinator_unclear_student,
    record_payment,
    reverse_payment,
    sync_presentations_from_registered,
    ensure_presentation_record,
    invalidate_cache,
    verify_supervisor,
    get_assigned_supervisor,
    get_students_for_supervisor,
    assign_supervisor,
    bulk_assign_supervisors,
    get_all_assignments,
    set_supervisor_passphrase,
    delete_student_record,
    sync_all_assignments,
    clean_presentations_sheet,
    get_all_registered_students,
    lookup_student_in_master,
    student_exists_v2,
    register_student_v2,
    verify_student_v2,
    update_student_password,
    migrate_students_sheet,
    is_portal_open,
    get_portal_message,
    get_presentation_fee,
    get_setting,
    set_setting,
    get_supervisors,
    get_supervisor_names,
    add_supervisor,
    update_supervisor_status,
    get_topic_proposals,
    decide_topic_proposal,
    get_available_topics,
    get_taken_topics,
    get_taken_topics_by_session,
    register_topic,
    is_student_registered,
    drop_registered_topic,
    find_student_record,
    set_panel_accepted,                     # <-- ADDED
)

PROGRAMMES = [
    'HND Software & Web Development',
    'HND Networking & Cloud Computing',
    'ND Computer Science',
]

@app.context_processor
def inject_globals():
    return {
        'current_year': datetime.now().year,
        'current_session': CURRENT_SESSION,
        'user_role': session.get('role', ''),
        'is_coordinator': session.get('role') == 'coordinator',
        'is_supervisor': session.get('role') == 'supervisor',
    }

# ── Decorators ────────────────────────────────────────────────────────────────
def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get('logged_in'):
            flash('Please log in first.', 'warning')
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated

def staff_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if session.get('role') not in ('coordinator', 'supervisor'):
            flash('Staff login required.', 'warning')
            return redirect(url_for('staff_login'))
        return f(*args, **kwargs)
    return decorated

def coordinator_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if session.get('role') != 'coordinator':
            flash('Project Coordinator access only.', 'danger')
            return redirect(url_for('coordinator_login'))
        return f(*args, **kwargs)
    return decorated

# ── Home ──────────────────────────────────────────────────────────────────────
@app.route('/')
def home():
    return render_template('index.html')

# ── Student Login ─────────────────────────────────────────────────────────────
@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        m = request.form.get('matric_number','').strip().upper()
        p = request.form.get('password','')
        if not m or not p:
            flash('Enter both matric number and password.', 'danger')
        else:
            # Verify password against Students sheet
            student = verify_student_v2(m, p)
            if not student:
                flash('Invalid matric number or password.', 'danger')
            else:
                # Check if student exists in AssignedSupervisors and has a supervisor
                identity = lookup_student_in_master(m)
                if not identity:
                    flash('Your record was not found in the assigned supervisors list. Please contact the Project Coordinator.', 'danger')
                    return redirect(url_for('login'))
                assigned_supervisor = identity.get('Assigned Supervisor', '').strip()
                if not assigned_supervisor:
                    flash('You have not been assigned a supervisor yet. Please contact the Project Coordinator before logging in.', 'warning')
                    return redirect(url_for('login'))
                # All checks passed – log the student in
                session.clear()
                session.update({
                    'logged_in': True,
                    'matric_number': student.get('Matric Number', m),
                    'student_name': student.get('Student Name', m),
                    'programme': student.get('Programme', ''),
                    'supervisor': assigned_supervisor,
                })
                flash(f"Welcome, {student.get('Student Name', m)}!", 'success')
                return redirect(url_for('submit_topic'))
    return render_template('login.html')

# ── Student Registration ──────────────────────────────────────────────────────
@app.route('/register', methods=['GET', 'POST'])
def register():
    prefill = {}
    if request.method == 'POST':
        action = request.form.get('action', 'validate')
        m = request.form.get('matric_number', '').strip().replace(' ', '').upper()
        email = request.form.get('email', '').strip()
        pw = request.form.get('password', '')
        cf = request.form.get('confirm_password', '')
        
        if action == 'validate':
            if not m:
                flash('Please enter your matric number.', 'danger')
            else:
                identity = lookup_student_in_master(m)
                if not identity:
                    flash(f'Matric number "{m}" was not found in the graduating students list.', 'danger')
                elif student_exists_v2(m):
                    flash('This matric number already has a portal account. Please log in instead.', 'warning')
                    return redirect(url_for('login'))
                else:
                    # Pre-fill all known details: Name, Matric, Programme, Supervisor
                    prefill = {
                        'Student Name': identity.get('Student Name', ''),
                        'Matric Number': identity.get('Matric Number', m),
                        'Programme': identity.get('Programme', ''),
                        'Assigned Supervisor': identity.get('Assigned Supervisor', '')
                    }
                    return render_template('register.html', prefill=prefill, step='create_account')
        
        elif action == 'create_account':
            if not all([m, email, pw, cf]):
                flash('All fields are required.', 'danger')
                prefill = lookup_student_in_master(m) or {}
            elif pw != cf:
                flash('Passwords do not match.', 'danger')
                prefill = lookup_student_in_master(m) or {}
            elif len(pw) < 6:
                flash('Password must be at least 6 characters.', 'danger')
                prefill = lookup_student_in_master(m) or {}
            elif student_exists_v2(m):
                flash('Account already exists. Please log in.', 'warning')
                return redirect(url_for('login'))
            else:
                ph = generate_password_hash(pw, method='scrypt')
                ok, msg = register_student_v2(m, email, ph)
                if ok:
                    # Fetch the student's name for the welcome message
                    identity = lookup_student_in_master(m)
                    student_name = identity.get('Student Name', '') if identity else ''
                    if student_name:
                        flash(f'Account created successfully. Welcome, {student_name}!', 'success')
                    else:
                        flash('Account created successfully. Please log in.', 'success')
                    return redirect(url_for('login'))
                else:
                    flash(f'Registration failed: {msg}', 'danger')
            return render_template('register.html', prefill=prefill, step='create_account')
    
    return render_template('register.html', prefill={}, step='validate')

# ── Student Logout ────────────────────────────────────────────────────────────
@app.route('/logout')
def logout():
    session.clear()
    flash('Logged out.', 'info')
    return redirect(url_for('login'))

# ── Submit Topic ──────────────────────────────────────────────────────────────

@app.route('/submit-topic', methods=['GET', 'POST'])
@login_required
def submit_topic():
    m = session['matric_number']
    prog = session.get('programme', '')
    just = session.pop('just_registered', False)
    assigned_supervisor = session.get('supervisor', '')

    # Optional: double-check that assigned_supervisor is not empty (though login already enforces)
    if not assigned_supervisor:
        flash('You have not been assigned a supervisor. Please contact the Project Coordinator.', 'danger')
        return redirect(url_for('logout'))

    if not is_portal_open():
        return render_template('portal_closed.html', message=get_portal_message())

    # Refresh programme from master sheet in case it changed (optional)
    identity = lookup_student_in_master(m)
    if identity:
        prog = identity.get('Programme', prog)
        session['programme'] = prog
        session['student_name'] = identity.get('Student Name', session.get('student_name', m))
        # supervisor should already be correct, but update if needed
        if identity.get('Assigned Supervisor'):
            session['supervisor'] = identity.get('Assigned Supervisor')
            assigned_supervisor = session['supervisor']

    topics = get_available_topics(prog)
    taken = is_student_registered(m)

    if request.method == 'POST':
        name = session['student_name']
        t = request.form['topic_title'].strip()
        # Use the supervisor from session (no form input)
        sup = assigned_supervisor
        if not all([name, m, prog, t, sup]):
            flash('All fields are required.', 'danger')
        elif taken:
            flash('Drop your existing topic first.', 'warning')
        elif register_topic(name, m, prog, t, sup):
            flash('Topic registered successfully!', 'success')
            session['just_registered'] = True
            # ── Email notification to supervisor ──────────────────────────
            try:
                supervisor_email = get_supervisor_email(sup)
                if supervisor_email:
                    notify_supervisor_topic_submitted(name, m, t, supervisor_email, sup)
            except Exception as e:
                # Log error but do not break the user experience
                app.logger.warning(f"Email notification failed: {e}")
            return redirect(url_for('submit_topic'))
        else:
            flash('Topic unavailable or already taken. Please choose another.', 'danger')

    return render_template('submit_topic.html',
                           topics=topics,
                           already_registered=taken,
                           just_registered=just,
                           assigned_supervisor=assigned_supervisor)

# ── View Available Topics ─────────────────────────────────────────────────────
@app.route('/view-topics')
def view_topics():
      
    # Allow only logged-in students or staff
    if not (session.get('logged_in') or session.get('role')):
        flash('Please log in to view available topics.', 'warning')
        return redirect(url_for('login'))
    
    prog = session.get('programme')
    if prog:
        topics = get_available_topics(prog)
    else:
        topics = {p: get_available_topics(p) for p in PROGRAMMES}
    return render_template('view_topics.html', topics=topics, programme=prog)

# ── Coordinator Login ─────────────────────────────────────────────────────────
@app.route('/coordinator-login', methods=['GET', 'POST'])
def coordinator_login():
    if request.method == 'POST':
        ph = request.form.get('passphrase', '').strip()
        coord_ph = get_coordinator_passphrase()
        if not ph:
            flash('Please enter a passphrase.', 'danger')
        elif coord_ph and ph == coord_ph:
            session.clear()
            session['role'] = 'coordinator'
            flash('Welcome, Project Coordinator.', 'success')
            return redirect(url_for('view_registered'))
        else:
            flash('Invalid coordinator passphrase.', 'danger')
    return render_template('coordinator_login.html')

# ── Staff Login (Supervisor) ───────────────────────────────────────────────────
@app.route('/staff-login', methods=['GET', 'POST'])
def staff_login():
    if request.method == 'POST':
        ph = request.form.get('passphrase', '').strip()
        if not ph:
            flash('Please enter a passphrase.', 'danger')
        else:
            # Only check against supervisor passphrases
            sup_name = verify_supervisor(ph)
            if sup_name:
                session.clear()
                session['role'] = 'supervisor'
                session['supervisor_name'] = sup_name
                flash(f'Welcome, {sup_name}.', 'success')
                return redirect(url_for('supervisor_clear_student_route'))
            else:
                flash('Invalid passphrase.', 'danger')
    return render_template('staff_login.html')

# ── Staff Logout ───────────────────────────────────────────────────────────────
@app.route('/staff-logout')
def staff_logout():
    role = session.get('role', '')
    session.pop('role', None)
    session.pop('supervisor_name', None)
    label = 'Project Coordinator' if role == 'coordinator' else 'Supervisor'
    flash(f'{label} logged out.', 'info')
    return redirect(url_for('staff_login'))

# ── View Registered Topics ────────────────────────────────────────────────────
@app.route('/view-registered')
@staff_required
def view_registered():
    view_mode = request.args.get('view', 'current')
    filter_prog = request.args.get('programme', '')
    role = session.get('role', '')
    if view_mode == 'current':
        raw = get_taken_topics()
    elif view_mode == 'past':
        raw = get_taken_topics_by_session(PREVIOUS_SESSION)
    else:  # all
        raw = get_taken_topics() + get_taken_topics_by_session(PREVIOUS_SESSION)
    registrations = {prog: [] for prog in PROGRAMMES}
    for rec in raw:
        prog_cell = rec.get('Programme', '').strip()
        if filter_prog and prog_cell.lower() != filter_prog.lower():
            continue
        matched = False
        for prog in PROGRAMMES:
            if prog.lower() == prog_cell.lower():
                registrations[prog].append({
                    'student_name': rec.get('Student Name', ''),
                    'matric_number': rec.get('Matric Number', ''),
                    'topic_title': rec.get('Topic Title', ''),
                    'supervisor': rec.get('Supervisor', ''),
                    'submission_date': rec.get('Submission Date', ''),
                    'session': rec.get('Session', CURRENT_SESSION),
                })
                matched = True
                break
    for prog in PROGRAMMES:
        registrations[prog].sort(key=lambda r: r['student_name'].lower())
    total = sum(len(v) for v in registrations.values())
    return render_template('view_registered.html', registrations=registrations,
                           programmes=PROGRAMMES, view_mode=view_mode,
                           filter_prog=filter_prog, current_session=CURRENT_SESSION,
                           total=total, role=role, is_coordinator=(role == 'coordinator'))

# ── Drop Topic ────────────────────────────────────────────────────────────────
@app.route('/drop-topic', methods=['POST'])
@login_required
def drop_topic():
    m = session['matric_number']
    prog = session.get('programme')
    if not prog:
        flash('Programme information missing. Please contact support.', 'danger')
        return redirect(url_for('submit_topic'))
    
    if drop_registered_topic(m, prog):
        flash('Topic dropped successfully.', 'success')
    else:
        flash('Could not drop topic. Please try again.', 'danger')
    return redirect(url_for('submit_topic'))

# ── Forgot Password ───────────────────────────────────────────────────────────
@app.route('/forgot-password', methods=['GET', 'POST'])
@app.route('/forgot_password', methods=['GET', 'POST'])
def forgot_password():
    if request.method == 'POST':
        m = request.form.get('matric_number','').strip().upper()
        npw = request.form.get('new_password','')
        cf = request.form.get('confirm_password','')
        if not all([m, npw, cf]):
            flash('All fields are required.', 'danger')
        elif npw != cf:
            flash('Passwords do not match.', 'danger')
        elif len(npw) < 6:
            flash('Password must be at least 6 characters.', 'danger')
        else:
            identity = lookup_student_in_master(m)
            if not identity:
                flash('Matric number not found in graduating students list.', 'danger')
            elif not student_exists_v2(m):
                flash('No portal account found. Please register first.', 'warning')
                return redirect(url_for('register'))
            else:
                hashed = generate_password_hash(npw, method='scrypt')
                if update_student_password(m, hashed):
                    flash('Password updated — please log in.', 'success')
                    return redirect(url_for('login'))
                else:
                    flash('Could not update password. Please try again.', 'danger')
        return redirect(url_for('forgot_password'))
    return render_template('forgot_password.html')

# ── Coordinator: Edit a student record ───────────────────────────────────────
@app.route('/admin/edit-student', methods=['GET', 'POST'])
@coordinator_required
def admin_edit_student():
    all_students = sorted(get_taken_topics(), key=lambda r: r.get('Student Name','').lower())
    sessions = [PREVIOUS_SESSION, CURRENT_SESSION]
    
    # Attach assigned supervisor (read-only)
    for student in all_students:
        matric = student.get('Matric Number', '')
        student['assigned_supervisor'] = get_assigned_supervisor(matric)
    
    if request.method == 'POST':
        orig_matric = request.form.get('orig_matric','').strip()
        new_name = request.form.get('student_name','').strip()
        new_matric = request.form.get('matric_number','').strip()
        new_prog = request.form.get('programme','').strip()
        new_topic = request.form.get('topic_title','').strip()
        new_super = request.form.get('supervisor','').strip()
        new_session = request.form.get('session','').strip()
        if not orig_matric:
            flash('No student selected.', 'danger')
        elif not all([new_name, new_matric, new_prog, new_topic, new_super, new_session]):
            flash('All fields are required.', 'danger')
        else:
            result = edit_student_record(orig_matric, new_name, new_matric, new_prog, new_topic, new_super, new_session)
            if result:
                flash(f'Record for {new_matric} updated successfully.', 'success')
            else:
                flash(f'Could not find {orig_matric} in TakenTopics.', 'warning')
        return redirect(url_for('admin_edit_student'))
    
    return render_template('admin_edit_student.html', students=all_students,
                           sessions=sessions, programmes=PROGRAMMES,
                           current_session=CURRENT_SESSION, previous_session=PREVIOUS_SESSION)


# ── Supervisor: Propose a topic or topic change ────────────────────────────────
@app.route('/supervisor/propose-topic', methods=['GET', 'POST'])
@staff_required
def supervisor_propose_topic():
    role = session.get('role')
    supervisor_name = session.get('supervisor_name', '')
    
    if request.method == 'POST':
        proposal_type = request.form.get('proposal_type','')
        if role == 'coordinator':
            proposer = request.form.get('proposer_name', 'Project Coordinator').strip()
        else:
            proposer = supervisor_name
        programme = request.form.get('programme','').strip()
        new_topic = request.form.get('new_topic','').strip()
        student_matric = request.form.get('student_matric','').strip()
        note = request.form.get('note','').strip()
        if not all([proposal_type, proposer, new_topic]):
            flash('Proposer name, type, and topic are required.', 'danger')
        else:
            ok = submit_topic_proposal(proposer, proposal_type, programme, new_topic, student_matric, note)
            if ok:
                # Send email notification to coordinator
                try:
                    from email_notify import notify_coordinator_proposal_submitted
                    notify_coordinator_proposal_submitted(proposer, proposal_type, new_topic, student_matric)
                except Exception as e:
                    app.logger.warning(f"Coordinator email notification failed: {e}")
                flash('Proposal submitted — awaiting Project Coordinator approval.', 'success')
            else:
                flash('Failed to submit proposal. Try again.', 'danger')
        return redirect(url_for('supervisor_propose_topic'))

    # ── GET: build student list based on role ──────────────────────────────
    all_registered = get_taken_topics()
    programmes = PROGRAMMES
    pending = get_topic_proposals(status='Pending') if role == 'coordinator' else []
    default_proposer = supervisor_name if role == 'supervisor' else 'Project Coordinator'

    if role == 'coordinator':
        students = sorted(all_registered, key=lambda r: r.get('Student Name','').lower())
    else:
        assigned_matrics = get_students_for_supervisor(supervisor_name)
        students = sorted(
            [s for s in all_registered if s.get('Matric Number', '').strip() in assigned_matrics],
            key=lambda r: r.get('Student Name','').lower()
        )
        if not students:
            flash('No students are currently assigned to you. Contact the coordinator.', 'info')

    return render_template('supervisor_propose_topic.html',
                           students=students,
                           programmes=programmes,
                           pending=pending,
                           is_coordinator=(role == 'coordinator'),
                           default_proposer=default_proposer)

# ── Coordinator: Review and approve/reject proposals ──────────────────────────
@app.route('/admin/review-proposals')
@coordinator_required
def admin_review_proposals():
    pending = get_topic_proposals(status='Pending')
    approved = get_topic_proposals(status='Approved')
    rejected = get_topic_proposals(status='Rejected')
    return render_template('admin_review_proposals.html', pending=pending,
                           approved=approved, rejected=rejected)

@app.route('/admin/decide-proposal', methods=['POST'])
@coordinator_required
def admin_decide_proposal():
    proposal_id = request.form.get('proposal_id','').strip()
    decision = request.form.get('decision','').strip()
    if not proposal_id or decision not in ('Approved', 'Rejected'):
        flash('Invalid request.', 'danger')
        return redirect(url_for('admin_review_proposals'))
    
    # Call the updated decide_topic_proposal that returns extra details
    ok, msg, proposer_name, proposer_email, proposal_type, new_topic = decide_topic_proposal(proposal_id, decision)
    
    # Send email to proposer if decision was made and email exists
    if ok and proposer_email:
        from email_notify import notify_proposer_decision
        notify_proposer_decision(proposer_email, proposer_name, proposal_type, new_topic, decision)
    
    flash(msg, 'success' if ok else 'warning')
    return redirect(url_for('admin_review_proposals'))

# ── Admin: Correct a student's session ───────────────────────────────────────
@app.route('/admin/update-session', methods=['GET', 'POST'])
@coordinator_required
def admin_update_session():
    all_students = get_taken_topics()
    sessions = [PREVIOUS_SESSION, CURRENT_SESSION]
    if request.method == 'POST':
        matric = request.form.get('matric_number','').strip()
        new_sess = request.form.get('session','').strip()
        if not matric or not new_sess:
            flash('Matric number and session are required.', 'danger')
        elif new_sess not in sessions:
            flash('Invalid session value.', 'danger')
        else:
            result = update_student_session(matric, new_sess)
            if result:
                flash(f'Session for {matric} updated to {new_sess}.', 'success')
            else:
                flash(f'Could not find {matric} in TakenTopics.', 'warning')
        return redirect(url_for('admin_update_session'))
    return render_template('admin_update_session.html',
                           students=sorted(all_students, key=lambda r: r.get('Student Name','').lower()),
                           sessions=sessions, current_session=CURRENT_SESSION,
                           previous_session=PREVIOUS_SESSION)

# ── Backfill Submission Dates ─────────────────────────────────────────────────
@app.route('/admin/backfill-dates')
@coordinator_required
def admin_backfill_dates():
    flash('Backfill not needed with session-specific tabs.', 'info')
    return redirect(url_for('view_registered'))

@app.route('/admin/backfill-sessions')
@coordinator_required
def admin_backfill_sessions():
    flash('Backfill not needed with session-specific tabs.', 'info')
    return redirect(url_for('view_registered'))

DEFAULT_PRESENTATION_FEE = 2000

# ── Supervisor: Clear a student for presentation ──────────────────────────────
@app.route('/supervisor/clear-student', methods=['GET', 'POST'])
@staff_required
def supervisor_clear_student_route():
    role = session.get('role', '')
    if request.method == 'POST':
        action = request.form.get('action', '').strip()
        matric = request.form.get('matric', '').strip()
        
        # Determine who is clearing the student
        if role == 'supervisor':
            cleared_by = session.get('supervisor_name', '')
        else:
            cleared_by = request.form.get('cleared_by', 'Project Coordinator').strip()
            if not cleared_by:
                cleared_by = 'Project Coordinator'
        
        if not matric:
            flash('No student selected.', 'danger')
        elif action == 'clear':
            if not cleared_by:
                cleared_by = session.get('supervisor_name', 'Unknown')
            ensure_presentation_record(matric)
            ok, msg = supervisor_clear_student(matric, cleared_by)
            if ok:
                # Notify student
                student_email = get_student_email(matric)
                if student_email:
                    student_name = get_student_name(matric)
                    if student_name:
                        notify_student_cleared(student_email, student_name, cleared_by)
                # Notify coordinator
                student_name = get_student_name(matric)
                if student_name:
                    notify_coordinator_cleared(student_name, matric, cleared_by)
            flash(msg, 'success' if ok else 'warning')
        elif action == 'unclear':
            ok, msg = coordinator_unclear_student(matric)
            flash(msg, 'success' if ok else 'warning')
        elif action == 'pay' and role == 'coordinator':
            fee = get_presentation_fee()
            ok, msg = record_payment(matric, fee, cleared_by)
            flash(msg, 'success' if ok else 'warning')
        elif action == 'unpay' and role == 'coordinator':
            ok, msg = reverse_payment(matric)
            flash(msg, 'success' if ok else 'warning')
        return redirect(url_for('supervisor_clear_student_route'))

    # GET request: build student list
    all_registered = get_taken_topics()
    assignments = get_all_assignments()
    assignment_map = {a.get("Matric Number", "").strip().lower(): a.get("Supervisor", "").strip() for a in assignments}
    
    if role == 'coordinator':
        students = sorted(all_registered, key=lambda r: r.get('Student Name', '').lower())
        my_name = 'Project Coordinator'
    else:
        sup_name = session.get('supervisor_name', '')
        sup_key = sup_name.lower()
        assigned_matrics = {a.get("Matric Number", "").strip().lower() for a in assignments if a.get("Supervisor", "").strip().lower() == sup_key}
        students = sorted([s for s in all_registered if s.get('Matric Number', '').strip().lower() in assigned_matrics],
                          key=lambda r: r.get('Student Name', '').lower())
        my_name = sup_name

    pres_records = {r['Matric Number'].strip().lower(): r for r in get_presentation_records(CURRENT_SESSION)}
    return render_template('supervisor_clear_student.html', students=students,
                           pres_records=pres_records, assignment_map=assignment_map,
                           current_session=CURRENT_SESSION, is_coordinator=(role == 'coordinator'),
                           supervisor_name=session.get('supervisor_name', ''),
                           default_name=my_name, presentation_fee=get_presentation_fee())

# ── Coordinator: Manage presentations ──────────────────────────────────────────
@app.route('/admin/manage-presentations', methods=['GET', 'POST'])
@coordinator_required
def admin_manage_presentations():
    if request.method == 'POST':
        action = request.form.get('action','')
        matric = request.form.get('matric','').strip()
        if action == 'sync':
            created = sync_presentations_from_registered()
            flash(f'Sync complete — {created} new presentation record(s) created.', 'success')
        elif action == 'pay':
            amount = request.form.get('amount', str(get_presentation_fee()))
            marked_by = request.form.get('marked_by', 'Project Coordinator').strip()
            try:
                amount = int(amount)
            except ValueError:
                amount = get_presentation_fee()
            ok, msg = record_payment(matric, amount, marked_by)
            flash(msg, 'success' if ok else 'warning')
        elif action == 'unpay':
            ok, msg = reverse_payment(matric)
            flash(msg, 'success' if ok else 'warning')
        elif action == 'clear':
            cleared_by = request.form.get('cleared_by', 'Project Coordinator').strip()
            ok, msg = supervisor_clear_student(matric, cleared_by)
            flash(msg, 'success' if ok else 'warning')
        elif action == 'unclear':
            ok, msg = coordinator_unclear_student(matric)
            flash(msg, 'success' if ok else 'warning')
        elif action == 'accept_panel':
            ok, msg = set_panel_accepted(matric, True, session.get('supervisor_name', 'Coordinator'))
            flash(msg, 'success' if ok else 'warning')
        return redirect(url_for('admin_manage_presentations'))

    registered = get_taken_topics()
    reg_matrics = {s.get('Matric Number','').strip().lower() for s in registered}
    all_pres = get_presentation_records(CURRENT_SESSION)
    pres_map = {r['Matric Number'].strip().lower(): r for r in all_pres}
    records = []
    for s in sorted(registered, key=lambda r: r.get('Student Name','').lower()):
        mkey = s.get('Matric Number','').strip().lower()
        pr = pres_map.get(mkey, {})
        panel_accepted = pr.get('Panel Accepted', '').strip().lower() == 'yes'
        records.append({
            'Student Name': s.get('Student Name',''),
            'Matric Number': s.get('Matric Number',''),
            'Programme': s.get('Programme',''),
            'Topic Title': s.get('Topic Title',''),
            'Supervisor': s.get('Supervisor',''),
            'Session': s.get('Session', CURRENT_SESSION),
            'Supervisor Cleared': pr.get('Supervisor Cleared','No'),
            'Cleared By': pr.get('Cleared By',''),
            'Cleared Date': pr.get('Cleared Date',''),
            'Payment Amount': pr.get('Payment Amount',''),
            'Payment Status': pr.get('Payment Status','Unpaid'),
            'Payment Date': pr.get('Payment Date',''),
            'Marked By': pr.get('Marked By',''),
            'eligible': (pr.get('Supervisor Cleared','').lower() == 'yes' and pr.get('Payment Status','').lower() == 'paid'),
            'panel_accepted': panel_accepted,
        })
    eligible = [r for r in records if r['eligible']]
    cleared_only = [r for r in records if r.get('Supervisor Cleared','').lower() == 'yes' and r.get('Payment Status','').lower() != 'paid']
    paid_only = [r for r in records if r.get('Payment Status','').lower() == 'paid' and r.get('Supervisor Cleared','').lower() != 'yes']
    neither = [r for r in records if r.get('Supervisor Cleared','').lower() != 'yes' and r.get('Payment Status','').lower() != 'paid']
    accepted = [r for r in records if r.get('panel_accepted', False)]
    not_synced = [r for r in all_pres if r.get('Matric Number','').strip().lower() not in reg_matrics and r.get('Session','') == CURRENT_SESSION]
    return render_template('admin_manage_presentations.html', 
                           eligible=eligible,
                           cleared_only=cleared_only, 
                           paid_only=paid_only,
                           neither=neither, 
                           accepted=accepted,
                           not_synced=not_synced, 
                           records=records,
                           current_session=CURRENT_SESSION, 
                           default_fee=get_presentation_fee())

# ── Coordinator: Printable presentation list (eligible = cleared + paid) ──────
@app.route('/admin/presentation-list')
@staff_required
def admin_presentation_list():
    status_filter = request.args.get('status', 'eligible')
    programme_filter = request.args.get('programme', '')
    role = session.get('role', '')
    supervisor_name = session.get('supervisor_name', '')

    # Get only current session records
    all_records = get_presentation_records(CURRENT_SESSION)

    # Supervisors: only see their assigned students
    if role == 'supervisor' and supervisor_name:
        assigned_matrics = get_students_for_supervisor(supervisor_name)
        all_records = [r for r in all_records if r.get('Matric Number', '').strip() in assigned_matrics]

    # Default for everyone is 'eligible' (cleared + paid)
    if status_filter == 'eligible':
        records = [r for r in all_records if r.get('eligible', False)]
    elif status_filter == 'cleared':
        records = [r for r in all_records if r.get('Supervisor Cleared', '').strip().lower() == 'yes']
    elif status_filter == 'paid':
        records = [r for r in all_records if r.get('Payment Status', '').strip().lower() == 'paid']
    elif status_filter == 'uncleared':
        records = [r for r in all_records if r.get('Supervisor Cleared', '').strip().lower() != 'yes']
    elif status_filter == 'unpaid':
        records = [r for r in all_records if r.get('Payment Status', '').strip().lower() != 'paid']
    else:
        records = all_records

    if programme_filter:
        records = [r for r in records if r.get('Programme', '').strip().lower() == programme_filter.lower()]

    grouped = {prog: [] for prog in PROGRAMMES}
    for r in records:
        prog = r.get('Programme', '')
        if prog in grouped:
            grouped[prog].append(r)

    return render_template('admin_presentation_list.html', grouped=grouped, records=records,
                           programmes=PROGRAMMES, status_filter=status_filter,
                           programme_filter=programme_filter, current_session=CURRENT_SESSION,
                           total=len(records), is_coordinator=(role == 'coordinator'),
                           generated_at=datetime.now().strftime("%d %B %Y, %I:%M %p"))

@app.route('/coordinator/accepted-topics')
@coordinator_required
def accepted_topics():
    """Show and export topics accepted by the panel (post-presentation)."""
    from flask import send_file, request
    import io, openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    programme_filter = request.args.get('programme', '')
    view_mode = request.args.get('view', 'html')

    records = get_presentation_records(CURRENT_SESSION)
    # Filter: Panel Accepted = Yes
    accepted = [r for r in records if r.get("Panel Accepted", "").strip().lower() == "yes"]

    if programme_filter:
        accepted = [r for r in accepted if r.get("Programme", "").lower() == programme_filter.lower()]

    if view_mode == 'excel':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Accepted Topics"
        headers = ["S/N", "Student Name", "Matric Number", "Programme", "Topic Title", "Supervisor", "Panel Accepted Date", "Cleared By", "Payment Date"]
        ws.append(headers)
        for sn, rec in enumerate(accepted, 1):
            ws.append([
                sn,
                rec.get("Student Name", ""),
                rec.get("Matric Number", ""),
                rec.get("Programme", ""),
                rec.get("Topic Title", ""),
                rec.get("Supervisor", ""),
                rec.get("Payment Date", ""),   # placeholder; can be changed later
                rec.get("Cleared By", ""),
                rec.get("Payment Date", "")
            ])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"Accepted_Topics_{CURRENT_SESSION}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        return send_file(buf, as_attachment=True, download_name=filename,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    return render_template('accepted_topics.html', accepted=accepted, programmes=PROGRAMMES,
                           programme_filter=programme_filter, current_session=CURRENT_SESSION,
                           generated_at=datetime.now().strftime("%d %B %Y, %I:%M %p"))

# ── Coordinator: Clean Presentations sheet ───────────────────────────────────
@app.route('/admin/clean-presentations', methods=['POST'])
@coordinator_required
def admin_clean_presentations():
    removed, fixed = clean_presentations_sheet()
    flash(f'Presentations cleaned: {removed} stale rows removed, {fixed} session values corrected.', 'success')
    return redirect(url_for('admin_manage_presentations'))

# ── Coordinator: Delete a student record ─────────────────────────────────────
@app.route('/admin/delete-student', methods=['POST'])
@coordinator_required
def admin_delete_student():
    matric = request.form.get('matric', '').strip()
    reason = request.form.get('reason', 'Duplicate entry removed by coordinator').strip()
    if not matric:
        flash('No matric number provided.', 'danger')
        return redirect(url_for('admin_edit_student'))
    ok, msg = delete_student_record(matric, reason)
    flash(msg, 'success' if ok else 'danger')
    return redirect(url_for('admin_edit_student'))

# ── Coordinator: Assign Supervisors to Students ──────────────────────────────
@app.route('/admin/assign-supervisors', methods=['GET', 'POST'])
@coordinator_required
def admin_assign_supervisors():
    filter_prog = request.args.get('programme', PROGRAMMES[0])
    if request.method == 'POST':
        action = request.form.get('action', '')
        prog_ctx = request.form.get('programme_ctx', filter_prog)
        if action == 'assign_one':
            matric = request.form.get('matric', '').strip()
            sup_name = request.form.get('supervisor', '').strip()
            student_name = request.form.get('student_name', '').strip()
            programme = request.form.get('programme_field', '').strip()
            if not matric or not sup_name:
                flash('Matric and supervisor are required.', 'danger')
            else:
                ok, msg = assign_supervisor(matric, student_name, programme, sup_name, 'Project Coordinator')
                flash(msg, 'success' if ok else 'warning')
                invalidate_cache()
        elif action == 'assign_bulk':
            assignments_to_make = []
            for key, val in request.form.items():
                if key.startswith('sup_') and val.strip():
                    matric = key[4:]
                    sup_name = val.strip()
                    sname = request.form.get(f'sname_{matric}', '').strip()
                    sprog = request.form.get(f'sprog_{matric}', '').strip()
                    assignments_to_make.append({'matric': matric, 'student_name': sname, 'programme': sprog, 'supervisor': sup_name})
            if assignments_to_make:
                ok_n, err_n, errors = bulk_assign_supervisors(assignments_to_make, 'Project Coordinator')
                flash(f'Saved {ok_n} assignment(s) for {prog_ctx}.', 'success' if err_n == 0 else 'warning')
                if errors:
                    flash(errors[0], 'danger')
            else:
                flash('No assignments to save — select supervisors first.', 'info')
        elif action == 'sync_takentopics':
            synced = sync_all_assignments()
            flash(f'Synced {synced} supervisor assignment(s) to registered topics.', 'success' if synced > 0 else 'info')
            invalidate_cache()
        return redirect(url_for('admin_assign_supervisors', programme=prog_ctx))

    all_reg = get_all_registered_students()
    all_students = [s for s in all_reg if s.get('Programme','').strip().lower() == filter_prog.strip().lower()]
    supervisors = get_supervisor_names()
    assignments = get_all_assignments()
    assignment_map = {a.get("Matric Number","").strip().lower(): a.get("Supervisor","").strip() for a in assignments}
    all_matrics = {s.get('Matric Number','').strip().lower() for s in all_reg}
    unassigned = len(all_matrics - set(assignment_map.keys()))
    return render_template('admin_assign_supervisors.html', students=all_students,
                           supervisors=supervisors, assignment_map=assignment_map,
                           current_session=CURRENT_SESSION, filter_prog=filter_prog,
                           programmes=PROGRAMMES, total_students=len(all_reg),
                           total_unassigned=unassigned)

# ── Coordinator: Migrate Students sheet to new format ────────────────────────
@app.route('/admin/migrate-students', methods=['POST'])
@coordinator_required
def admin_migrate_students():
    ok, msg = migrate_students_sheet()
    flash(msg, 'success' if ok else 'danger')
    return redirect(url_for('admin_settings'))

# ── Coordinator: Portal Settings ─────────────────────────────────────────────
@app.route('/admin/settings', methods=['GET', 'POST'])
@coordinator_required
def admin_settings():
    if request.method == 'POST':
        action = request.form.get('action', '')
        if action == 'toggle_portal':
            currently_open = is_portal_open()
            set_setting('portal_open', 'false' if currently_open else 'true')
            flash(f'Portal {"closed" if currently_open else "opened"} successfully.', 'success')
        elif action == 'set_message':
            msg = request.form.get('portal_message', '').strip()
            if msg:
                set_setting('portal_message', msg)
                flash('Closed-portal message updated.', 'success')
        elif action == 'set_fee':
            try:
                fee = int(request.form.get('fee', 2000))
                set_setting('presentation_fee', fee)
                flash(f'Presentation fee set to ₦{fee:,}.', 'success')
            except ValueError:
                flash('Invalid fee amount.', 'danger')
        return redirect(url_for('admin_settings'))
    return render_template('admin_settings.html', portal_open=is_portal_open(),
                           portal_message=get_portal_message(), presentation_fee=get_presentation_fee())

# ── Coordinator: Manage Supervisors ──────────────────────────────────────────
@app.route('/admin/supervisors', methods=['GET', 'POST'])
@coordinator_required
def admin_supervisors():
    if request.method == 'POST':
        action = request.form.get('action', '')
        if action == 'add':
            full_name = request.form.get('full_name', '').strip()
            short_name = request.form.get('short_name', '').strip()
            department = request.form.get('department', '').strip()
            if not full_name:
                flash('Full name is required.', 'danger')
            else:
                ok, msg = add_supervisor(full_name, short_name, department)
                flash(msg, 'success' if ok else 'warning')
        elif action == 'deactivate':
            name = request.form.get('name', '').strip()
            ok, msg = update_supervisor_status(name, 'No')
            flash(msg, 'success' if ok else 'warning')
        elif action == 'activate':
            name = request.form.get('name', '').strip()
            ok, msg = update_supervisor_status(name, 'Yes')
            flash(msg, 'success' if ok else 'warning')
        elif action == 'set_passphrase':
            name = request.form.get('name', '').strip()
            passphrase = request.form.get('passphrase', '').strip()
            if not passphrase:
                flash('Passphrase cannot be empty.', 'danger')
            else:
                ok, msg = set_supervisor_passphrase(name, passphrase)
                flash(msg, 'success' if ok else 'warning')
        return redirect(url_for('admin_supervisors'))
    all_supervisors = get_supervisors(active_only=False)
    return render_template('admin_supervisors.html', supervisors=all_supervisors)

# ── Download registered topics as Excel ──────────────────────────────────────
@app.route('/download/registered-topics')
@app.route('/download_registered_topics')
@staff_required
def download_registered_topics():
    import io, openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from flask import send_file
    view_mode = request.args.get('view', 'current')
    filter_prog = request.args.get('programme', '')
    if view_mode == 'current':
        raw = get_taken_topics()
    elif view_mode == 'past':
        raw = get_taken_topics_by_session(PREVIOUS_SESSION)
    else:
        raw = get_taken_topics() + get_taken_topics_by_session(PREVIOUS_SESSION)
    if filter_prog:
        raw = [r for r in raw if r.get('Programme','').lower() == filter_prog.lower()]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registered Topics"
    headers = ["S/N", "Student Name", "Matric Number", "Programme", "Topic Title", "Supervisor", "Date Submitted", "Session"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    row_num = 2
    sn = 1
    for rec in raw:
        ws.cell(row=row_num, column=1, value=sn)
        ws.cell(row=row_num, column=2, value=rec.get('Student Name',''))
        ws.cell(row=row_num, column=3, value=rec.get('Matric Number',''))
        ws.cell(row=row_num, column=4, value=rec.get('Programme',''))
        ws.cell(row=row_num, column=5, value=rec.get('Topic Title',''))
        ws.cell(row=row_num, column=6, value=rec.get('Supervisor',''))
        ws.cell(row=row_num, column=7, value=rec.get('Submission Date','')[:10])
        ws.cell(row=row_num, column=8, value=rec.get('Session',''))
        row_num += 1
        sn += 1
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"Registered_Topics_{view_mode}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 5000)), debug=False)